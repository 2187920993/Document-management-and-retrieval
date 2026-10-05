import json
import re
from pathlib import Path

from .config import STORAGE_DIR, TEXT_EXTENSIONS
from .db import db, is_postgres
from .embeddings import embed


def extract_text(path: Path, extension: str) -> str:
    if extension in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace")
    if extension == ".docx":
        try:
            from docx import Document as WordDocument
            document = WordDocument(str(path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(line for line in paragraphs if line.strip())
        except Exception:
            return ""
    if extension == ".pptx":
        try:
            from pptx import Presentation
            presentation = Presentation(str(path))
            lines = []
            for slide_index, slide in enumerate(presentation.slides, start=1):
                slide_lines = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        slide_lines.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
                    if getattr(shape, "has_table", False):
                        slide_lines.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                if slide_lines:
                    lines.append(f"第 {slide_index} 页\n" + "\n".join(line for line in slide_lines if line.strip()))
            return "\n\n".join(lines)
        except Exception:
            return ""
    if extension == ".xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        rows.append(" | ".join(values))
                if rows:
                    lines.append(f"工作表：{sheet.title}\n" + "\n".join(rows))
            return "\n\n".join(lines)
        except Exception:
            return ""
    if extension == ".xls":
        try:
            import xlrd
            workbook = xlrd.open_workbook(str(path), on_demand=True)
            lines = []
            for sheet in workbook.sheets():
                rows = []
                for row_index in range(sheet.nrows):
                    values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
                    if values:
                        rows.append(" | ".join(values))
                if rows:
                    lines.append(f"工作表：{sheet.name}\n" + "\n".join(rows))
            return "\n\n".join(lines)
        except Exception:
            return ""
    if extension == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        except Exception:
            return ""
    return ""


def split_chunks(text: str, max_chars: int = 900, overlap: int = 120) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 1 <= max_chars:
            current = f"{current}\n{paragraph}".strip()
        else:
            if current:
                chunks.append(current)
            current = (current[-overlap:] + "\n" + paragraph).strip() if overlap else paragraph
    if current:
        chunks.append(current)
    return chunks or ([text[:max_chars]] if text.strip() else [])


def _row_value(row, key):
    return row[key] if isinstance(row, dict) else row[key]


def index_document(document_id: str) -> None:
    with db() as conn:
        row = conn.execute("SELECT id, storage_key, original_name, extension, content_version FROM documents WHERE id = %s" if is_postgres() else "SELECT id, storage_key, original_name, extension, content_version FROM documents WHERE id = ?", (document_id,)).fetchone()
        if not row:
            return
        path = STORAGE_DIR / _row_value(row, "storage_key")
        text = extract_text(path, _row_value(row, "extension")) if path.exists() else ""
        chunks = split_chunks(text)
        source_name = _row_value(row, "original_name") or ""
        # Keep the source name in the persisted vector input so semantic search
        # can find a document even when its title is more informative than its
        # extracted body (and when a PDF has little usable text).
        if not chunks and source_name:
            chunks = [f"文件名：{source_name}"]
        conn.execute("DELETE FROM document_chunks WHERE document_id = %s" if is_postgres() else "DELETE FROM document_chunks WHERE document_id = ?", (document_id,))
        for index, chunk in enumerate(chunks):
            vector = json.dumps(embed(f"文件名：{source_name}\n{chunk}"))
            if is_postgres():
                conn.execute("INSERT INTO document_chunks (document_id, content_version, chunk_index, text, vector) VALUES (%s, %s, %s, %s, %s::jsonb)", (document_id, _row_value(row, "content_version"), index, chunk, vector))
            else:
                conn.execute("INSERT INTO document_chunks (document_id, content_version, chunk_index, text, vector) VALUES (?, ?, ?, ?, ?)", (document_id, _row_value(row, "content_version"), index, chunk, vector))
        status = "ready" if chunks else ("not_supported" if _row_value(row, "extension") in {".pdf", ".docx", ".ppt", ".pptx", ".xls", ".xlsx"} else "failed")
        error = None if chunks else "未提取到可检索正文"
        conn.execute("UPDATE documents SET index_status = %s, index_error = %s WHERE id = %s" if is_postgres() else "UPDATE documents SET index_status = ?, index_error = ? WHERE id = ?", (status, error, document_id))
